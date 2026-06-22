"""
Generador de presentación PPT para el proyecto Grupo E.
Ejecutar: python presentacion/presentacion.py
"""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
import os

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

BLUE = RGBColor(0x1A, 0x23, 0x7E)
DARK = RGBColor(0x1E, 0x1E, 0x1E)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
ORANGE = RGBColor(0xFF, 0x57, 0x22)
GRAY = RGBColor(0x66, 0x66, 0x66)
LIGHT_BG = RGBColor(0xF5, 0xF5, 0xF5)


def add_bg(slide, color=DARK):
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_text_box(slide, left, top, width, height, text, font_size=18,
                 bold=False, color=WHITE, alignment=PP_ALIGN.LEFT):
    txBox = slide.shapes.add_textbox(Inches(left), Inches(top),
                                      Inches(width), Inches(height))
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.bold = bold
    p.font.color.rgb = color
    p.alignment = alignment
    return tf


def add_bullet_slide(slide, left, top, width, height, items, font_size=16, color=WHITE):
    txBox = slide.shapes.add_textbox(Inches(left), Inches(top),
                                      Inches(width), Inches(height))
    tf = txBox.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.text = f"• {item}"
        p.font.size = Pt(font_size)
        p.font.color.rgb = color
        p.space_after = Pt(8)
    return tf


# --- Slide 1: Portada ---
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, DARK)
add_text_box(slide, 1, 1.5, 11, 1, "Navegación Urbana", 48, True, WHITE, PP_ALIGN.CENTER)
add_text_box(slide, 1, 2.8, 11, 0.8, "con Caminos más Cortos", 40, True, ORANGE, PP_ALIGN.CENTER)
add_text_box(slide, 1, 4.2, 11, 0.6, "Dijkstra vs A* sobre la Red Vial de Lima", 22, False, GRAY, PP_ALIGN.CENTER)
add_text_box(slide, 1, 5.5, 11, 0.5, "Grupo E — Análisis y Diseño de Algoritmos 2026", 18, False, WHITE, PP_ALIGN.CENTER)
add_text_box(slide, 1, 6.2, 11, 0.4, "Universidad ESAN", 16, False, GRAY, PP_ALIGN.CENTER)

# --- Slide 2: Problema ---
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, DARK)
add_text_box(slide, 0.8, 0.5, 11, 0.8, "Problema", 36, True, ORANGE)
add_bullet_slide(slide, 0.8, 1.6, 11, 5, [
    "Encontrar la ruta más corta entre dos puntos en una ciudad",
    "Aplicación real: navegadores GPS, logística, transporte urbano",
    "Datos reales: red vial de Lima desde OpenStreetMap",
    "Algoritmos: Dijkstra (clásico) vs A* (con heurística)",
    "Métrica: distancia en metros sobre vías reales",
], 18, WHITE)

# --- Slide 3: Modelado ---
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, DARK)
add_text_box(slide, 0.8, 0.5, 11, 0.8, "Modelado del Grafo", 36, True, ORANGE)
add_bullet_slide(slide, 0.8, 1.6, 11, 5, [
    "Grafo G = (V, E) dirigido y ponderado",
    "V = intersecciones viales (nodos)",
    "E = segmentos de vía (aristas con peso = longitud en metros)",
    "Fuente: OpenStreetMap vía OSMnx (network_type='drive')",
    "Lima: ~100,000 nodos, ~250,000 aristas",
], 18, WHITE)

# --- Slide 4: Algoritmos ---
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, DARK)
add_text_box(slide, 0.8, 0.5, 11, 0.8, "Algoritmos Implementados", 36, True, ORANGE)

add_text_box(slide, 0.8, 1.6, 5.5, 0.5, "Dijkstra", 28, True, WHITE)
add_bullet_slide(slide, 0.8, 2.3, 5.5, 3, [
    "Expande en orden de distancia",
    "Garantiza optimalidad",
    "Heap binario como cola prioridad",
    "O((V+E) log V)",
], 16, WHITE)

add_text_box(slide, 6.8, 1.6, 5.5, 0.5, "A*", 28, True, ORANGE)
add_bullet_slide(slide, 6.8, 2.3, 5.5, 3, [
    "Guía la búsqueda con heurística",
    "Heurística Haversine (distancia línea recta)",
    "Admisible → garantiza optimalidad",
    "Explora menos nodos en la práctica",
], 16, WHITE)

# --- Slide 5: Heap ---
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, DARK)
add_text_box(slide, 0.8, 0.5, 11, 0.8, "Heap Binario", 36, True, ORANGE)
add_bullet_slide(slide, 0.8, 1.6, 11, 4, [
    "Cola de prioridad implementada desde cero",
    "push(p, x): O(log n) — insertar o actualizar",
    "pop(): O(log n) — extraer el mínimo",
    "Estructura de árbol binario completa",
], 18, WHITE)

# --- Slide 6: Heurística ---
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, DARK)
add_text_box(slide, 0.8, 0.5, 11, 0.8, "Heurística Haversine", 36, True, ORANGE)
add_bullet_slide(slide, 0.8, 1.6, 11, 5, [
    "Distancia del gran círculo sobre la esfera terrestre",
    "d = 2R · arcsin(√(sin²(Δφ/2) + cos φ₁·cos φ₂·sin²(Δλ/2)))",
    "R = 6371 km (radio de la Tierra)",
    "Admisible: nunca sobrestima el costo real",
    "Consistente: garantiza optimalidad de A*",
], 18, WHITE)

# --- Slide 7: Dataset ---
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, DARK)
add_text_box(slide, 0.8, 0.5, 11, 0.8, "Dataset: Red Vial de Lima", 36, True, ORANGE)
add_bullet_slide(slide, 0.8, 1.6, 11, 4, [
    "OpenStreetMap — datos colaborativos mundiales",
    "OSMnx: descarga directa y conversión a NetworkX",
    "Alternativa: archivo .osm.pbf local (Perú completo)",
    "Filtro: vías transitables para automóvil ('drive')",
], 18, WHITE)

# --- Slide 8: Resultados ---
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, DARK)
add_text_box(slide, 0.8, 0.5, 11, 0.8, "Resultados Experimentales", 36, True, ORANGE)

metrics = [
    ("Métrica", "Dijkstra", "A*", "Mejora"),
    ("Tiempo (s)", "t_d", "t_a", "t_d / t_a ×"),
    ("Nodos explorados", "n_d", "n_a", "n_d / n_a ×"),
    ("Distancia (m)", "d", "d", "Igual"),
]
y = 1.6
for row in metrics:
    for i, cell in enumerate(row):
        c = ORANGE if row == metrics[0] else WHITE
        b = True if row == metrics[0] else False
        add_text_box(slide, 0.8 + i * 3, y, 2.8, 0.5, cell, 18, b, c, PP_ALIGN.CENTER)
    y += 0.7

add_text_box(slide, 0.8, 5.0, 11, 0.8, "Ambos algoritmos encuentran la misma ruta óptima. A* explora menos nodos y es más rápido.",
             16, False, GRAY, PP_ALIGN.CENTER)

# --- Slide 9: Conclusiones ---
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, DARK)
add_text_box(slide, 0.8, 0.5, 11, 0.8, "Conclusiones", 36, True, ORANGE)
add_bullet_slide(slide, 0.8, 1.6, 11, 5, [
    "A* supera a Dijkstra en nodos explorados y tiempo",
    "Heurística Haversine es adecuada para redes urbanas",
    "Heap binario implementado correctamente",
    "OSMnx facilita trabajo con datos viales reales",
    "Código reproducible y disponible públicamente",
], 18, WHITE)

# --- Slide 10: Demo ---
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, DARK)
add_text_box(slide, 0.8, 0.5, 11, 0.8, "Demo en Vivo", 44, True, ORANGE, PP_ALIGN.CENTER)
add_text_box(slide, 0.8, 3.0, 11, 1, "Ejecución del buscador de rutas\nsobre la red vial de Lima",
             28, False, WHITE, PP_ALIGN.CENTER)
add_text_box(slide, 0.8, 5.0, 11, 0.6, "python demo.py [--interactive]", 20, False, GRAY, PP_ALIGN.CENTER)

# --- Slide 11: Gracias ---
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, DARK)
add_text_box(slide, 0.8, 2.5, 11, 1, "Gracias", 48, True, WHITE, PP_ALIGN.CENTER)
add_text_box(slide, 0.8, 4.0, 11, 0.8, "¿Preguntas?", 32, False, ORANGE, PP_ALIGN.CENTER)
add_text_box(slide, 0.8, 5.5, 11, 0.5, "Grupo E — ADA 2026 — Universidad ESAN", 18, False, GRAY, PP_ALIGN.CENTER)

output_path = os.path.join(os.path.dirname(__file__), "Presentacion_Grupo_E.pptx")
prs.save(output_path)
print(f"Presentación guardada en: {output_path}")
